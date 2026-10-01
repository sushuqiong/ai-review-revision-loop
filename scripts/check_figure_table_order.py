#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""figure/table ORDER-COMPLIANCE probe — run before EVERY delivery.

Why this exists: on v22 the manuscript was submitted to Frontiers with tables numbered
1..7 while the text first cited them 1,7,6,2,3,4,5, and supplementary tables numbered
S1..S16 while first cited in a near-random order. Integrity checks all passed, because
none of them asserted ORDER. This script is that missing assertion.

Checks four classes on the MAIN manuscript (citation first-mention order must ascend)
and caption ARRANGEMENT order on the SUPPLEMENTARY file (physical order must ascend).
Optionally verifies embedded figure resolution + byte-identity with Figures/.

Exit code 0 = compliant, 1 = violations found.

Usage:
  python check_figure_table_order.py Manuscript.docx ["Supplementary Materials.docx"] [--figures Figures]
"""
import io
import os
import re
import sys
import zipfile
import hashlib

# --- patterns: main tables must NOT swallow "Table S13"; supplementary use the S form ---
P_MAIN_T = r'\bTables?\s+(\d+)(?![\dS])'
P_SUPP_T = r'\bTables?\s+S(\d+)'
P_MAIN_F = r'\bFigures?\s+(\d+)(?![\dS])'
P_SUPP_F = r'\bFigures?\s+S(\d+)'
# caption = paragraph that STARTS with "Table 3." / "Table S13." / "Figure 2." / "Figure S1."
RE_CAP = re.compile(r'^(Table|Figure)\s+S?(\d+)\.')
# range refs, en-dash \u2013 or hyphen
RE_RANGE_S = re.compile(r'\bTable\s+S(\d+)\s*[\u2013-]\s*S?(\d+)')
RE_RANGE_T = re.compile(r'\bTable\s+(\d+)\s*[\u2013-]\s*(\d+)')


def load(path):
    from docx import Document
    d = Document(path)
    paras = [p.text for p in d.paragraphs]
    cells = [c.text for t in d.tables for tr in t.rows for c in tr.cells]
    return d, paras, cells


def first_mention(paras, pat):
    """{number: index of first paragraph mentioning it}"""
    out = {}
    for i, t in enumerate(paras):
        for m in re.finditer(pat, t):
            out.setdefault(int(m.group(1)), i)
    return out


def order_report(label, firsts):
    seq = [k for k, _ in sorted(firsts.items(), key=lambda kv: kv[1])]
    ok = seq == sorted(seq)
    print("  %-11s first-mention order %s  %s" % (label, seq, "OK" if ok else "*** NOT ASCENDING ***"))
    if not ok:
        print("  %-11s   expected ascending: %s" % ("", sorted(seq)))
    return ok


def caption_order(doc):
    """captions in physical document order, split main/supplementary"""
    from docx.oxml.ns import qn
    mt, mf, st, sf = [], [], [], []
    for el in doc.element.body.iterchildren():
        if el.tag != qn('w:p'):
            continue
        t = "".join(n.text or "" for n in el.iter(qn('w:t'))).strip()
        m = RE_CAP.match(t)
        if not m:
            continue
        kind, num = m.group(1), int(m.group(2))
        is_supp = bool(re.match(r'^(Table|Figure)\s+S\d+\.', t))
        if is_supp:
            (st if kind == "Table" else sf).append(num)
        else:
            (mt if kind == "Table" else mf).append(num)
    return mt, mf, st, sf


def trailing_run(seq):
    """the end-matter attachment block = trailing run that restarts at 1"""
    out = []
    for x in reversed(seq):
        if out and x > out[-1]:
            break
        out.append(x)
    return list(reversed(out))


def out_of_range(text, hi_main, hi_supp):
    bad = []
    for m in RE_RANGE_S.finditer(text):
        for g in m.groups():
            if not (1 <= int(g) <= hi_supp):
                bad.append(m.group(0))
    for m in RE_RANGE_T.finditer(text):
        for g in m.groups():
            if not (1 <= int(g) <= hi_main):
                bad.append(m.group(0))
    for pat, hi in ((P_SUPP_T, hi_supp), (P_MAIN_T, hi_main), (P_SUPP_F, hi_supp), (P_MAIN_F, hi_main)):
        for m in re.finditer(pat, text):
            if not (1 <= int(m.group(1)) <= hi):
                bad.append(m.group(0))
    return bad


def check_figures(figdir, pairs):
    """pairs: [(docx_path, [(media_name, figure_filename), ...]), ...]"""
    bad = 0
    for docx_path, ps in pairs:
        z = zipfile.ZipFile(docx_path)
        for media, fname in ps:
            try:
                data = z.read("word/media/" + media)
            except KeyError:
                print("  MISSING %s in %s" % (media, os.path.basename(docx_path)))
                bad += 1
                continue
            ref = os.path.join(figdir, fname)
            if not os.path.exists(ref):
                print("  MISSING reference file %s" % ref)
                bad += 1
                continue
            same = hashlib.md5(data).hexdigest() == hashlib.md5(io.open(ref, "rb").read()).hexdigest()
            try:
                from PIL import Image
                w = Image.open(io.BytesIO(data)).size[0]
            except Exception:
                w = -1
            hires = w >= 2000
            flag = "" if (same and hires) else "   <-- PROBLEM"
            print("  %-24s %-14s %s px  md5-match=%s%s" % (os.path.basename(docx_path)[:24], fname, w, same, flag))
            if not (same and hires):
                bad += 1
    return bad


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    figdir = None
    if "--figures" in sys.argv:
        figdir = sys.argv[sys.argv.index("--figures") + 1]
    if not args:
        print(__doc__)
        return 2
    ms = args[0]
    supp = args[1] if len(args) > 1 else None
    failures = 0

    print("=" * 72)
    print(" figure/table ORDER compliance: %s" % os.path.basename(ms))
    print("=" * 72)

    dms, pm, cm = load(ms)
    print("\n[1] citation first-mention order (MAIN manuscript)")
    for label, pat in (("Table", P_MAIN_T), ("Table S", P_SUPP_T), ("Figure", P_MAIN_F), ("Figure S", P_SUPP_F)):
        fm = first_mention(pm, pat)
        if fm and not order_report(label, fm):
            failures += 1

    print("\n[2] caption arrangement (MAIN manuscript end matter)")
    mt, mf, st, sf = caption_order(dms)
    tr = trailing_run(mt)
    if tr:
        exp = list(range(1, len(tr) + 1))
        ok = tr == exp
        print("  Table blocks %s  %s" % (tr, "OK" if ok else "*** NOT 1..N ***"))
        if not ok:
            failures += 1

    hi_t = max([x for x in mt if x] or [7])
    hi_s = max([x for x in st if x] or [16])

    if supp:
        print("\n[3] SUPPLEMENTARY caption arrangement (physical order)")
        ds, ps, cs = load(supp)
        _, _, st2, sf2 = caption_order(ds)
        for lbl, arr in (("Table S", st2), ("Figure S", sf2)):
            ok = arr == sorted(arr) and arr == list(range(1, len(arr) + 1))
            print("  %-9s %s  %s" % (lbl, arr, "OK" if ok else "*** NOT 1..N ASCENDING ***"))
            if not ok:
                failures += 1

    print("\n[4] out-of-range references")
    alltext = "\n".join(pm + cm)
    if supp:
        alltext += "\n" + "\n".join(ps + cs)
    bad = out_of_range(alltext, hi_t, hi_s)
    print("  %s" % ("OK (none)" if not bad else "*** %s ***" % sorted(set(bad))))
    if bad:
        failures += 1

    if figdir:
        print("\n[5] embedded figures vs %s" % figdir)
        pairs = [(ms, [("image2.png", "Fig1.png"), ("image3.png", "Fig2.png"), ("image4.png", "Fig3.png")])]
        if supp:
            pairs.append((supp, [("image1.png", "FigureS1.png"), ("image2.png", "FigureS2.png"), ("image3.png", "FigureS3.png")]))
        failures += check_figures(figdir, pairs)

    print("\n" + "=" * 72)
    print(" RESULT: %d violation group(s)" % failures)
    print(" NOTE: 'my checklist passed' != 'the manuscript has no compliance problems'.")
    print("       Never report 'ready to submit' on the strength of this script alone.")
    print("=" * 72)
    return 1 if failures else 0


if __name__ == "__main__":
    sys.exit(main())
