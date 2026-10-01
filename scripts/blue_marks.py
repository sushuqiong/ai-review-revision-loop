# -*- coding: utf-8 -*-
"""把 <O>…<C> 标记包住的片段染蓝（用于生成"改动标蓝"核对版 docx）。

用法：
    from blue_marks import apply_blue_marks, set_blue, O, C, BLUE

    # a) 整段是新增/重写的 → 直接给段落上色
    set_blue(p)                       # 或用生成器里的 set_text(p, txt, blue=True)

    # b) 长段落里只新增一句 → 用标记包住那一句
    txt = old_prefix + O + "新增的那句。" + C + old_suffix
    ...
    apply_blue_marks(doc)             # save 之前调用一次

产出两份交付件：
    NO_BLUE=1 python gen.py           # 先跑黑版 → 复制到 投稿用_去蓝标版本/
    python gen.py                     # 再跑蓝版（用户点名的路径）

设计要点见 references/change-highlighting-for-review.md：
  * 不要用 diff 自动标蓝（引用转上标后 token 序列处处错位，会把大片"其实没变"的文字标蓝）
  * 交付前必验：核对版蓝标数 == 改动处数、去蓝标版 == 0、两者正文逐字相同、无残留标记字符
"""
import copy
import os

from docx.shared import RGBColor

# 私用区字符（U+E000/U+E001），正常论文正文里绝不会出现
O, C = "\ue000", "\ue001"
BLUE = RGBColor(0x00, 0x00, 0xFF)

# 设 NO_BLUE=1 时只剥离标记、不上色 → 得到可直接投稿的干净件
NO_BLUE = bool(os.environ.get("NO_BLUE"))

__all__ = ["apply_blue_marks", "set_blue", "O", "C", "BLUE", "NO_BLUE"]


def set_blue(paragraph, colour=None):
    """把整段染蓝（保留原格式，只改颜色）。NO_BLUE 时不动。"""
    if NO_BLUE:
        return 0
    colour = BLUE if colour is None else colour
    n = 0
    for run in paragraph.runs:
        if run.text.strip():
            run.font.color.rgb = colour
            n += 1
    return n


def apply_blue_marks(doc, colour=None):
    """把标记对之间的文本染蓝、并剥离标记字符。返回被染蓝的片段数。

    逐 run 处理：把 run 按标记切成若干片段，克隆原 run 元素以继承全部格式，
    奇数片段（位于一对标记之间）上色。不会破坏已有的上标 run —— 上标 run 里
    没有标记，原样保留。
    """
    colour = BLUE if colour is None else colour
    if NO_BLUE:
        colour = None
    n = 0
    for p in doc.paragraphs:
        if O not in p.text and C not in p.text:
            continue
        for run in list(p.runs):
            t = run.text
            if O not in t and C not in t:
                continue
            parts = [x for x in t.replace(C, O).split(O)]
            template = copy.deepcopy(run._element)
            parent = run._element.getparent()
            anchor = run._element
            anchor.addprevious(copy.deepcopy(template))
            prev = anchor.getprevious()
            used = False
            for i, seg in enumerate(parts):
                if seg == "":
                    continue
                nr = prev if not used else copy.deepcopy(template)
                if used:
                    prev.addnext(nr)
                from docx.text.run import Run
                R = Run(nr, p)
                R.text = seg
                if i % 2 == 1 and colour is not None:
                    R.font.color.rgb = colour
                    n += 1
                prev = nr
                used = True
            parent.remove(anchor)
    return n


def audit(blue_path, clean_path):
    """交付前自检：返回 (蓝标段数, 干净版蓝标段数, 正文是否逐字相同, 残留标记数)。"""
    from docx import Document

    def blue_count(path):
        d = Document(path)
        k = 0
        for p in d.paragraphs:
            for r in p.runs:
                try:
                    c = r.font.color
                    if c is not None and c.rgb is not None and str(c.rgb) == "0000FF" and r.text.strip():
                        k += 1
                        break
                except Exception:
                    pass
        return k

    tb = [p.text for p in Document(blue_path).paragraphs]
    tc = [p.text for p in Document(clean_path).paragraphs]
    stray = sum(p.text.count(O) + p.text.count(C) for p in Document(blue_path).paragraphs)
    return blue_count(blue_path), blue_count(clean_path), tb == tc, stray
