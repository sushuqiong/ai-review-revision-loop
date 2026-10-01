#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
qa_revised_manuscript.py — 返修稿交付前的程序化自检（可直接重跑）

用法：
    python qa_revised_manuscript.py <revised_manuscript.docx> [title_page.docx] [--quiet]

检查项（每项独立打印 OK / !!）：
  1. 正文无残留 [n] 方括号引用
  2. 上标引用编号的「首次出现顺序」严格 == 1..N（不是"个数差不多"）
  3. 参考文献条目编号严格 1..N；被引未列 / 已列未被引 == 0
  4. 摘要字数（结构化 4 段空白分词）== 标题页申报值
  5. 正文字数（Introduction → References 空白分词）== 标题页申报值
  6. 手稿内图注条数（Fig. N）
  7. 英文文档无中文字符泄漏
  8. 术语统一：Supplemental / Supplementary 只留一个
  9. 图表引用带完整前缀（不出现裸 "Table S1"）

退出码：0 = 全过；1 = 有失败项。

注意（这三个坑本次都踩过，本脚本已修正）：
  * 段落是 list，list.count(x) 数的是"等于 x 的元素"，不是子串出现次数 → 一律 join 成整串再 count
  * 参考文献正则用 r"^\\d+\\.\\s"，不要加 [A-Z][a-z]（会漏掉 "8. GBD 2019 Hepatitis B Collaborators."）
  * 图注正则用 r"^Fig\\.\\s?\\d"，容忍 "Fig. 1" 里的空格
"""
import os
import re
import sys

LAB_DEFAULT = ("Introduction and Objectives", "Materials and methods", "Results:", "Conclusions:")
CITE = re.compile(r"\[\s*(\d+(?:\s*[\u2013\-,]\s*\d+)*)\s*\]")
SUP_ONLY = re.compile(r"^[\d,\s\u2013\-]+$")

FAILS = []


def ok(cond, label, extra=""):
    if not cond:
        FAILS.append(label)
    print("   [%s] %s %s" % ("OK " if cond else "!! ", label, extra))
    return cond


def wc(text):
    """空白分词——与作者标题页申报口径一致（本次已用提交版 245 词验证过）。"""
    return len(re.sub(r"\s+", " ", text).strip().split())


def paras(path):
    from docx import Document
    return Document(path)


def expand_ranges(nums_text):
    out = []
    for part in re.split(r"[,\s]+", nums_text.strip()):
        if not part:
            continue
        if re.search(r"[\u2013\-]", part):
            a, b = re.split(r"[\u2013\-]", part)
            if a.isdigit() and b.isdigit():
                out.extend(range(int(a), int(b) + 1))
        elif part.isdigit():
            out.append(int(part))
    return out


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    quiet = "--quiet" in sys.argv
    if not args:
        print(__doc__)
        return 2
    ms_path = args[0]
    tp_path = args[1] if len(args) > 1 else None

    doc = paras(ms_path)
    ps = doc.paragraphs
    texts = [p.text for p in ps]

    try:
        i_refs = next(i for i, t in enumerate(texts) if t.strip() == "References")
    except StopIteration:
        print("!! 找不到 'References' 段——请确认这是手稿正文文件")
        return 1
    i_intro = next((i for i, t in enumerate(texts) if re.match(r"^1\.\s+\S", t.strip())), 0)

    print("=" * 92)
    print("QA:", os.path.basename(ms_path), "| 段落 %d | References@%d" % (len(ps), i_refs))
    print("=" * 92)

    # ---- 1. 残留方括号引用
    body_raw = "\n".join(texts[:i_refs])
    print("\n[1] 正文残留 [n] 方括号引用")
    ok(len(CITE.findall(body_raw)) == 0, "正文无未转换的 [n]")

    # ---- 2. 上标编号首现序
    print("\n[2] 上标引用：首现顺序与完整性")
    seq = []
    for p in ps[:i_refs]:
        for r in p.runs:
            if r.font.superscript and SUP_ONLY.match(r.text.strip()):
                seq += expand_ranges(r.text)
    first = []
    for n in seq:
        if n not in first:
            first.append(n)
    n_ref = len(first)
    print("   上标引用点数 %d | 去重后 %d" % (len(seq), n_ref))
    if n_ref:
        expected = list(range(1, n_ref + 1))
        missing = sorted(set(range(1, n_ref + 1)) - set(first))
        ok(first == expected, "编号按首现顺序且无跳号（1..%d）" % n_ref,
           ("缺: %s" % missing) if missing else "")
    ok(all(a <= b for a, b in zip([0] + first, first)), "首现序单调递增（无回引提前出现）")

    # ---- 3. 参考文献表
    print("\n[3] 参考文献表")
    refs = [t.strip() for t in texts[i_refs + 1:] if re.match(r"^\d+\.\s", t.strip())]
    nums = [int(t.split(".")[0]) for t in refs]
    ok(nums == list(range(1, len(refs) + 1)), "条目编号严格 1..N", "%d 条" % len(refs))
    if refs:
        ok(sorted(set(first)) == list(range(1, len(refs) + 1)),
           "被引集合 == 列表编号集合（无被引未列 / 已列未引）",
           "已列未被引: %s" % (sorted(set(range(1, len(refs) + 1)) - set(first)) or "无"))

    # ---- 4/5. 字数
    print("\n[4][5] 字数（空白分词，与标题页同口径）")
    lab = tuple(l for l in LAB_DEFAULT if any(t.strip().startswith(l) for t in texts[:i_intro]))
    if lab:
        abs_txt = " ".join(t for t in texts[:i_intro] if t.strip().startswith(lab))
    else:                                    # 回退：ABSTRACT 之后、首个章节标题之前
        a0 = next((i for i, t in enumerate(texts[:i_intro]) if t.strip().upper() == "ABSTRACT"), 0)
        abs_txt = " ".join(t for t in texts[a0 + 1:i_intro] if t.strip() and not t.strip().lower().startswith("keywords"))
    n_abs = wc(abs_txt)
    n_main = wc(" ".join(texts[i_intro:i_refs]))
    print("   摘要 %d 词 | 正文 %d 词" % (n_abs, n_main))
    if tp_path and os.path.exists(tp_path):
        tpt = "\n".join(p.text for p in paras(tp_path).paragraphs)
        m_a = re.search(r"Abstract:?\s*([\d,]+)\s*words", tpt, re.I)
        m_m = re.search(r"Main text[^:]*:\s*([\d,]+)\s*words", tpt, re.I)
        ok(bool(m_a) and int(m_a.group(1).replace(",", "")) == n_abs,
           "标题页摘要字数 == 实测", "标题页 %s / 实测 %d" % (m_a.group(1) if m_a else "缺失", n_abs))
        ok(bool(m_m) and int(m_m.group(1).replace(",", "")) == n_main,
           "标题页正文字数 == 实测", "标题页 %s / 实测 %d" % (m_m.group(1) if m_m else "缺失", n_main))
    else:
        print("   (未提供标题页，跳过字数比对)")

    # ---- 6. 手稿内图注
    print("\n[6] 手稿内图注")
    figleg = [t for t in texts if re.match(r"^Fig\.\s?\d", t.strip())]
    ok(bool(figleg), "存在手稿内图注区", "%d 条" % len(figleg))

    # ---- 7. 中文泄漏
    print("\n[7] 英文文档无中文字符")
    ok(not re.search(r"[\u4e00-\u9fff]", body_raw + "\n".join(texts[i_refs:])), "正文/参考文献无中文")

    # ---- 8. 术语统一
    print("\n[8] 术语统一 Supplementary / Supplemental")
    all_txt = "\n".join(texts)
    n_sup = len(re.findall(r"Supplemental\b", all_txt))
    ok(n_sup == 0, "不出现 'Supplemental'（统一用 Supplementary）", "出现 %d 次" % n_sup)

    # ---- 9. 图表引用前缀
    print("\n[9] 补充材料引用带完整前缀")
    bare = re.findall(r"(?<!Supplementary )(?<!supplementary )(?<!and )\bTables?\sS\d", body_raw)
    ok(len(bare) == 0, "不出现裸 'Table S1/S2'（应写 Supplementary Table S1）",
       ("发现: %s" % set(bare)) if bare else "")

    print("\n" + "=" * 92)
    print("结果：%d 项失败" % len(FAILS))
    for f in FAILS:
        print("   !!", f)
    print("=" * 92)
    return 1 if FAILS else 0


if __name__ == "__main__":
    sys.exit(main())
